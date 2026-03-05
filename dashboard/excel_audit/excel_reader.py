"""Excel reader for audit purposes.

Wraps ExcelParserService, groups parsed Deal objects by period,
and provides period-name filtering.
"""

from __future__ import annotations

import asyncio
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

from loguru import logger

_PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
_SRC = str(_PROJECT_ROOT / "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from application.excel_parser.parser import ExcelParserService  # noqa: E402


def get_available_periods(file_path: str) -> list[str]:
    """Return sheet names (period names) from the Excel file.

    Args:
        file_path: Path to the Excel workbook.

    Returns:
        list[str]: Ordered list of sheet names (whitespace stripped).
    """
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(file_path, read_only=True, data_only=True)
    # Strip whitespace — some sheets may have trailing spaces in their names
    sheets = [s.strip() for s in wb.sheetnames]
    wb.close()
    logger.info("Available periods in file: {}", sheets)
    return sheets


def parse_all_periods(file_path: str) -> dict[str, list[Any]]:
    """Parse Excel file and group Deal objects by period name.

    Uses ExcelParserService in async context via asyncio.run().
    Period name is formed as "{period_month} {period_year}" which matches
    the Excel sheet name format.

    Args:
        file_path: Path to the Excel workbook.

    Returns:
        dict[str, list[Deal]]: Mapping period_name → list of Deal objects.
    """
    logger.info("Parsing Excel file: {}", file_path)
    service = ExcelParserService(use_cache_only=None)
    result = asyncio.run(service.parse_file(file_path, None))

    grouped: dict[str, list] = defaultdict(list)
    for deal in result.deals:
        period_name = f"{deal.period_month} {deal.period_year}"
        grouped[period_name].append(deal)

    logger.info(
        "Parsed {} deals across {} period(s): {}",
        len(result.deals),
        len(grouped),
        list(grouped.keys()),
    )
    if result.stats.errors:
        for err in result.stats.errors:
            logger.warning("Parser warning: {}", err)

    return dict(grouped)


def filter_periods(
    all_periods: dict[str, list[Any]],
    period_filter: list[str] | None,
) -> dict[str, list[Any]]:
    """Filter period map by a list of requested period names.

    Args:
        all_periods: Full period → deals mapping.
        period_filter: List of period names to keep; None means keep all.

    Returns:
        dict[str, list[Deal]]: Filtered mapping.
    """
    if not period_filter:
        return all_periods

    missing = [p for p in period_filter if p not in all_periods]
    if missing:
        logger.warning("Requested periods not found in file: {}", missing)

    return {k: v for k, v in all_periods.items() if k in period_filter}
