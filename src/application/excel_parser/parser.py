"""
Excel Parser Application Service.

Refactored version of the original excel_parser.py using domain models and modern architecture.
"""

from __future__ import annotations

import hashlib
import os
import re
from decimal import Decimal
from enum import Enum
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from loguru import logger
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from xlcalculator import Evaluator, ModelCompiler

from domain.builders.deal_builder import DealBuilder
from domain.exceptions import SyncFileError
from domain.models import Deal, DealItem
from domain.value_objects import Money, Money5, Period, SignedMoney, SignedMoney5, Status

from .models import ParseResult, ParseStats
from .precalc import refresh_excel_cache_if_enabled

_DOTENV_LOADED: bool = False


def _load_env_if_needed() -> None:
    """Load configuration from .env file once per process."""
    global _DOTENV_LOADED

    if _DOTENV_LOADED:
        return

    if os.getenv("APP_DISABLE_DOTENV") == "1":
        logger.debug("dotenv loading disabled via APP_DISABLE_DOTENV=1")
        _DOTENV_LOADED = True
        return

    env_file = os.getenv("ENV_FILE", "config.env")
    env_path = Path(env_file).resolve()

    try:
        loaded = load_dotenv(dotenv_path=str(env_path), override=False)
        logger.debug(
            "dotenv loaded: file=%s loaded=%s override=%s",
            env_path,
            loaded,
            False,
        )
    except Exception as exc:
        logger.debug("dotenv load failed for %s: %s", env_path, str(exc))

    _DOTENV_LOADED = True


_load_env_if_needed()


class CacheMode(str, Enum):
    CACHE_ONLY = "cache_only"
    NO_CACHE = "no_cache"
    HYBRID = "hybrid"


class ExcelParserService:
    """Excel parser application service."""

    def __init__(self, use_cache_only: bool | None = None) -> None:
        self.stats = ParseStats()
        self.cache_mode = self._resolve_cache_mode(use_cache_only)
        logger.info(f"ExcelParserService cache mode: {self.cache_mode.value}")

    def _resolve_cache_mode(self, use_cache_only: bool | None) -> CacheMode:
        explicit_mode = os.getenv("EXCEL_PARSER_CACHE_MODE")

        if use_cache_only is not None:
            return CacheMode.CACHE_ONLY if use_cache_only else CacheMode.HYBRID

        if explicit_mode:
            normalized = explicit_mode.strip().lower()
            mode_map = {
                "cache_only": CacheMode.CACHE_ONLY,
                "no_cache": CacheMode.NO_CACHE,
                "hybrid": CacheMode.HYBRID,
            }
            if normalized in mode_map:
                return mode_map[normalized]
            logger.warning(
                "Unknown EXCEL_PARSER_CACHE_MODE='%s', falling back to 'hybrid'",
                explicit_mode,
            )

        return CacheMode.HYBRID

    async def parse_file(self, file_path: str, sync_session: object | None) -> ParseResult:
        """
        Parse Excel file and return structured result.

        Args:
            file_path: Path to Excel file
            sync_session: Associated sync session

        Returns:
            ParseResult with deals and statistics

        Raises:
            SyncFileError: If file cannot be read or processed
        """
        logger.info(f"Starting Excel file parsing: {file_path}")

        try:
            # Validate file existence and get metadata
            file_info = self._get_file_info(file_path)

            # Read Excel sheets
            sheets_data = self._read_excel_file(file_path)
            self.stats.total_sheets = len(sheets_data)

            all_deals = []

            # Process each sheet
            for sheet_name, df in sheets_data.items():
                logger.info(f"Processing sheet: {sheet_name}")
                try:
                    deals = await self._parse_sheet(df, sheet_name.strip())
                    if deals:  # Only count sheets that actually produced deals
                        all_deals.extend(deals)
                        self.stats.processed_sheets += 1
                        logger.info(f"✅ Sheet '{sheet_name}': found {len(deals)} deals")
                    else:
                        logger.info(f"📋 Sheet '{sheet_name}': skipped (no valid data)")
                except Exception as e:
                    self.stats.failed_sheets += 1
                    error_msg = f"Error processing sheet '{sheet_name}': {str(e)}"
                    self.stats.errors.append(error_msg)
                    logger.error(error_msg)
                    continue

            # Update sync session statistics
            if sync_session is not None and hasattr(sync_session, "stats"):
                sync_session.stats.total_deals = self.stats.total_deals
                sync_session.stats.processed_deals = self.stats.processed_deals
                sync_session.stats.failed_deals = self.stats.failed_deals
                sync_session.stats.total_items = self.stats.total_items
                sync_session.stats.processed_items = self.stats.processed_items
                sync_session.stats.failed_items = self.stats.failed_items
                sync_session.stats.errors.extend(self.stats.errors)
                sync_session.stats.warnings.extend(self.stats.warnings)

            # Create result
            result = ParseResult(
                deals=all_deals,
                stats=self.stats,
                sync_session=sync_session,
                file_path=file_path,
                file_size=file_info["size"],
                file_hash=file_info["hash"],
            )

            logger.info(f"Parsing completed. Total deals: {len(all_deals)}")
            self._log_final_stats()

            return result

        except Exception as e:
            # If it's already a SyncFileError, just re-raise it
            if isinstance(e, SyncFileError):
                raise

            error_msg = f"Critical error parsing file '{file_path}': {str(e)}"
            logger.error(error_msg)
            self.stats.errors.append(error_msg)
            raise SyncFileError(file_path, error_msg) from e

    def _get_file_info(self, file_path: str) -> dict:
        """Get file metadata for validation."""
        path = Path(file_path)

        if not path.exists():
            raise SyncFileError(file_path, "File does not exist")

        if not path.is_file():
            raise SyncFileError(file_path, "Path is not a file")

        # Calculate file hash
        with open(file_path, "rb") as f:
            file_hash = hashlib.md5(f.read()).hexdigest()

        return {"size": path.stat().st_size, "hash": file_hash}

    def _evaluate_simple_formula(self, formula: str, worksheet, current_row: int) -> float | None:
        """
        Evaluate simple Excel formulas.

        Supports:
        - Binary operations: =C5*E5, =C5+D5, =C5-D5, =C5/E5
        - Formulas with parentheses: =(E5-D5)*C5, =(A1+B1)/C1

        Args:
            formula: Formula string (e.g., '=C5*E5' or '=(E5-D5)*C5')
            worksheet: openpyxl worksheet
            current_row: Current row number for context

        Returns:
            Calculated value or None if cannot evaluate
        """
        try:
            def get_cell_value(cell_addr: str):
                col_match = re.match(r'([A-Z]+)(\d+)', cell_addr)
                if not col_match:
                    return None
                col_letter = col_match.group(1)
                row_num = int(col_match.group(2))
                col_idx = column_index_from_string(col_letter)
                cell = worksheet.cell(row_num, col_idx)
                value = cell.value
                try:
                    return float(value) if value is not None else None
                except (ValueError, TypeError):
                    return None

            # Remove = sign and spaces
            formula = formula.strip().lstrip('=').replace(' ', '')

            # Try pattern: =(CELL1-CELL2)*CELL3 or =(CELL1+CELL2)*CELL3
            pattern1 = r'\(([A-Z]+\d+)([\+\-])([A-Z]+\d+)\)\*([A-Z]+\d+)'
            match = re.match(pattern1, formula)
            if match:
                val1 = get_cell_value(match.group(1))
                op1 = match.group(2)
                val2 = get_cell_value(match.group(3))
                val3 = get_cell_value(match.group(4))

                if val1 is not None and val2 is not None and val3 is not None:
                    if op1 == '+':
                        return (val1 + val2) * val3
                    elif op1 == '-':
                        return (val1 - val2) * val3

            # Try pattern: =(CELL1-CELL2)/CELL3 or =(CELL1+CELL2)/CELL3
            pattern2 = r'\(([A-Z]+\d+)([\+\-])([A-Z]+\d+)\)/([A-Z]+\d+)'
            match = re.match(pattern2, formula)
            if match:
                val1 = get_cell_value(match.group(1))
                op1 = match.group(2)
                val2 = get_cell_value(match.group(3))
                val3 = get_cell_value(match.group(4))

                if val1 is not None and val2 is not None and val3 is not None and val3 != 0:
                    if op1 == '+':
                        return (val1 + val2) / val3
                    elif op1 == '-':
                        return (val1 - val2) / val3

            # Try simple binary operations: CELL1 OPERATOR CELL2
            pattern3 = r'([A-Z]+\d+)([\*\+\-\/])([A-Z]+\d+)'
            match = re.match(pattern3, formula)
            if match:
                val1 = get_cell_value(match.group(1))
                operator = match.group(2)
                val2 = get_cell_value(match.group(3))

                if val1 is not None and val2 is not None:
                    if operator == '*':
                        return val1 * val2
                    elif operator == '+':
                        return val1 + val2
                    elif operator == '-':
                        return val1 - val2
                    elif operator == '/' and val2 != 0:
                        return val1 / val2

            return None

        except Exception:
            return None

    def _manual_sum_evaluation(
        self, formula: str, evaluator, sheet_name: str, logger, worksheet=None
    ) -> float:
        """
        Manually evaluate SUM formula when xlcalculator fails.

        xlcalculator has known issues with SUM of formula ranges.
        This method extracts the range and manually sums each cell by evaluating
        their formulas directly.

        Args:
            formula: Excel formula string (e.g., '=SUM(F5:F16)')
            evaluator: xlcalculator Evaluator instance
            sheet_name: Sheet name
            logger: Logger instance
            worksheet: openpyxl worksheet (optional, for direct cell access)

        Returns:
            Calculated sum value
        """
        try:
            # Extract range from SUM formula (e.g., 'F5:F16' from '=SUM(F5:F16)')
            match = re.search(r'SUM\(([A-Z]+\d+):([A-Z]+\d+)\)', formula, re.IGNORECASE)
            if not match:
                logger.debug(f"Could not parse SUM range from formula: {formula}")
                return 0

            start_cell = match.group(1)  # e.g., 'F5'
            end_cell = match.group(2)    # e.g., 'F16'

            # Parse cell coordinates
            start_col = re.match(r'([A-Z]+)', start_cell).group(1)
            start_row = int(re.match(r'[A-Z]+(\d+)', start_cell).group(1))
            end_row = int(re.match(r'[A-Z]+(\d+)', end_cell).group(1))

            # Calculate sum manually by evaluating each cell's formula
            total = 0
            evaluated_count = 0

            if worksheet:
                # Use worksheet to get cell values/formulas directly
                col_idx = column_index_from_string(start_col)
                for row_num in range(start_row, end_row + 1):
                    cell = worksheet.cell(row_num, col_idx)
                    cell_value = cell.value

                    # If cell has a formula, try to evaluate it manually
                    if cell_value and isinstance(cell_value, str) and cell_value.startswith('='):
                        # Try simple formulas like =C5*E5
                        value = self._evaluate_simple_formula(cell_value, worksheet, row_num)
                        if value is not None and isinstance(value, (int, float)):
                            total += value
                            evaluated_count += 1
                    # If it's a direct value
                    elif isinstance(cell_value, (int, float)):
                        total += cell_value
                        evaluated_count += 1

            logger.debug(
                f"Manual SUM evaluation: {formula} = {total} "
                f"(evaluated {evaluated_count}/{end_row - start_row + 1} cells)"
            )
            return total

        except Exception as e:
            logger.warning(f"Manual SUM evaluation failed for {formula}: {str(e)}")
            return 0

    @staticmethod
    def _detect_formula_columns(worksheet) -> set[int]:
        """Identify columns containing formula cells to optimize per-cell checks."""
        formula_columns: set[int] = set()
        try:
            max_rows_to_scan = worksheet.max_row or 0
            if max_rows_to_scan == 0:
                return formula_columns

            max_rows_to_scan = min(max_rows_to_scan, 200)

            for row in worksheet.iter_rows(min_row=1, max_row=max_rows_to_scan):
                for col_idx, cell in enumerate(row, start=1):
                    value = cell.value
                    if isinstance(value, str) and value.startswith('='):
                        formula_columns.add(col_idx)
                if worksheet.max_column and len(formula_columns) == worksheet.max_column:
                    break
        except Exception:
            return formula_columns

        return formula_columns

    def _read_excel_file(self, file_path: str) -> dict:
        """Read Excel file and return sheets data with calculated formulas."""
        try:
            logger.info("Loading Excel file with formula calculation support...")

            original_path = Path(file_path)
            effective_path = refresh_excel_cache_if_enabled(original_path)
            if effective_path != original_path:
                logger.info("Using precalculated Excel copy: %s", effective_path)

            workbook_source = str(effective_path)

            # Load workbook with formulas (data_only=False to see formulas)
            workbook = load_workbook(filename=workbook_source, data_only=False)
            # Load workbook with cached values (data_only=True) for fallback
            workbook_values = None
            if self.cache_mode != CacheMode.NO_CACHE:
                workbook_values = load_workbook(
                    filename=workbook_source,
                    data_only=True,
                    read_only=True,
                )

            # Initialize xlcalculator to evaluate formulas (hybrid mode only)
            evaluator = None
            if self.cache_mode == CacheMode.HYBRID:
                try:
                    compiler = ModelCompiler()
                    model = compiler.read_and_parse_archive(workbook_source)
                    evaluator = Evaluator(model)
                    logger.info("Formula evaluator initialized successfully")
                except Exception as e:
                    logger.warning(f"Could not initialize formula evaluator: {str(e)}")
                    logger.warning("Falling back to reading cached values only")
                    evaluator = None

            logger.debug(
                "Excel parser mode settings: cache_enabled=%s evaluator_enabled=%s",
                bool(workbook_values),
                bool(evaluator),
            )

            sheets_data: dict[str, pd.DataFrame] = {}

            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                values_ws = (
                    workbook_values[sheet_name]
                    if workbook_values is not None
                    else None
                )
                rows = []

                # Detect header row and total_* columns by header name
                header_row_idx: int | None = None
                total_cols: set[int] = set()
                try:
                    max_scan = min(10, worksheet.max_row or 10)
                    for i, hdr_row in enumerate(worksheet.iter_rows(min_row=1, max_row=max_scan), start=1):
                        try:
                            if any(
                                (val is not None) and ("Клиент" in str(val))
                                for val in (c.value for c in hdr_row)
                            ):
                                header_row_idx = i
                                break
                        except Exception:
                            continue
                    if header_row_idx is not None:
                        header_values: list[str | None]
                        if values_ws is not None:
                            header_values = list(
                                values_ws.iter_rows(
                                    min_row=header_row_idx,
                                    max_row=header_row_idx,
                                    values_only=True,
                                )
                            )[0]
                        else:
                            header_rows = list(
                                worksheet.iter_rows(
                                    min_row=header_row_idx,
                                    max_row=header_row_idx,
                                    values_only=True,
                                )
                            )
                            header_values = list(header_rows[0]) if header_rows else []

                        for j, hval in enumerate(header_values, start=1):
                            if hval is None:
                                continue
                            name = str(hval).strip().lower()
                            if (
                                ("выруч" in name)
                                or ("марж" in name)
                                or ("закуп" in name)
                                or ("себестоим" in name)
                            ):
                                total_cols.add(j)
                except Exception:
                    # Non-fatal: if detection fails, we will not treat any column specially
                    header_row_idx = None

                formula_columns = self._detect_formula_columns(worksheet)

                iter_values = values_ws.iter_rows(values_only=True) if values_ws is not None else None
                for row_idx, row in enumerate(worksheet.iter_rows(), start=1):
                    value_row = next(iter_values) if iter_values is not None else None
                    row_values = []
                    for col_idx, cell in enumerate(row, start=1):
                        cached_value = None
                        if value_row is not None and col_idx - 1 < len(value_row):
                            cached_value = value_row[col_idx - 1]
                        cell_value = cell.value
                        has_formula_column = col_idx in formula_columns
                        is_formula = (
                            has_formula_column
                            and isinstance(cell_value, str)
                            and cell_value.startswith('=')
                        )
                        is_total_column = col_idx in total_cols

                        if self.cache_mode == CacheMode.CACHE_ONLY:
                            if is_formula:
                                row_values.append(cached_value)
                            else:
                                row_values.append(cell_value)
                            continue

                        if self.cache_mode == CacheMode.NO_CACHE:
                            if is_formula:
                                row_values.append(None)
                            else:
                                row_values.append(cell_value)
                            continue

                        should_skip_cache = False

                        if is_total_column:
                            if cached_value is not None:
                                row_values.append(cached_value)
                                continue

                            if evaluator and is_formula:
                                try:
                                    col_letter = get_column_letter(col_idx)
                                    cell_address = f"{sheet_name}!{col_letter}{row_idx}"
                                    calculated_value = evaluator.evaluate(cell_address)
                                    if (
                                        (calculated_value == 0 or calculated_value is None)
                                        and "SUM(" in cell_value.upper()
                                    ):
                                        calculated_value = self._manual_sum_evaluation(
                                            cell_value, evaluator, sheet_name, logger, worksheet
                                        )
                                    row_values.append(calculated_value)
                                except Exception as exc:
                                    logger.debug(
                                        "Formula evaluation failed for %s [%s,%s]: %s",
                                        sheet_name,
                                        row_idx,
                                        col_idx,
                                        str(exc),
                                    )
                                    row_values.append(None)
                            else:
                                if cached_value is not None:
                                    row_values.append(cached_value)
                                else:
                                    row_values.append(None)
                        else:
                            if evaluator and is_formula:
                                try:
                                    col_letter = get_column_letter(col_idx)
                                    cell_address = f"{sheet_name}!{col_letter}{row_idx}"
                                    calculated_value = evaluator.evaluate(cell_address)
                                    if (
                                        (calculated_value == 0 or calculated_value is None)
                                        and "SUM(" in cell_value.upper()
                                    ):
                                        calculated_value = self._manual_sum_evaluation(
                                            cell_value, evaluator, sheet_name, logger, worksheet
                                        )
                                    if calculated_value is None and cached_value is not None:
                                        row_values.append(cached_value)
                                    else:
                                        row_values.append(calculated_value)
                                except Exception:
                                    if cached_value is not None:
                                        row_values.append(cached_value)
                                    else:
                                        row_values.append(None)
                            else:
                                if is_formula:
                                    if cached_value is not None:
                                        row_values.append(cached_value)
                                    else:
                                        row_values.append(None)
                                else:
                                    # For price columns: col_idx 4=purchase_price, col_idx 7=margin (1-based Excel indexing)
                                    # Preserve as Decimal to avoid precision loss when pandas converts to float64
                                    # Note: col_idx starts from 1 (Excel column numbering)
                                    if col_idx in (4, 7):  # purchase_price (col 4), margin (col 7)
                                        if cell_value is not None:
                                            # Convert numeric values to Decimal to preserve precision
                                            if isinstance(cell_value, (int, float)):
                                                row_values.append(Decimal(str(cell_value)))
                                            elif isinstance(cell_value, Decimal):
                                                row_values.append(cell_value)
                                            else:
                                                try:
                                                    row_values.append(Decimal(str(cell_value)))
                                                except (ValueError, TypeError):
                                                    row_values.append(cell_value)
                                        elif cached_value is not None:
                                            if isinstance(cached_value, (int, float)):
                                                row_values.append(Decimal(str(cached_value)))
                                            elif isinstance(cached_value, Decimal):
                                                row_values.append(cached_value)
                                            else:
                                                try:
                                                    row_values.append(Decimal(str(cached_value)))
                                                except (ValueError, TypeError):
                                                    row_values.append(cached_value)
                                        else:
                                            row_values.append(None)
                                    else:
                                        if cell_value is not None:
                                            row_values.append(cell_value)
                                        elif cached_value is not None:
                                            row_values.append(cached_value)
                                        else:
                                            row_values.append(None)

                    rows.append(tuple(row_values))

                if rows:
                    df = pd.DataFrame(rows)
                else:
                    df = pd.DataFrame()

                sheets_data[sheet_name] = df

            workbook.close()
            if workbook_values is not None:
                try:
                    workbook_values.close()
                except Exception:
                    pass
            return sheets_data

        except Exception as e:
            raise SyncFileError(
                file_path, f"Cannot read Excel file: {str(e)}"
            ) from e

    async def _parse_sheet(self, df: pd.DataFrame, sheet_name: str) -> list[Deal]:
        """Parse individual Excel sheet."""
        if df.empty:
            logger.info(f"📋 Sheet '{sheet_name}' is empty - skipping")
            return []

        # Validate sheet name and extract period FIRST
        try:
            period = Period.from_sheet_name(sheet_name)
        except Exception as e:
            logger.info(f"📋 Sheet '{sheet_name}' has invalid period format - skipping: {str(e)}")
            return []

        # Find header row
        header_row = self._find_header_row(df)
        if header_row is None:
            logger.info(f"📋 Sheet '{sheet_name}' has no valid header row - skipping")
            return []

        if len(df) <= header_row:
            logger.info(f"📋 Sheet '{sheet_name}' has insufficient rows - skipping")
            return []

        logger.info(f"Sheet '{sheet_name}': header row at {header_row}")

        try:
            # Set headers and data
            df.columns = df.iloc[header_row]
            data_df = df.iloc[header_row + 1 :].reset_index(drop=True)
        except Exception as e:
            logger.info(f"📋 Sheet '{sheet_name}' has invalid headers - skipping: {str(e)}")
            return []

        # Parse deals
        return await self._parse_deals(data_df, period)

    def _find_header_row(self, df: pd.DataFrame) -> int | None:
        """Find row containing headers."""
        try:
            for i in range(min(5, len(df))):
                try:
                    row = df.iloc[i]
                    for val in row:
                        if pd.notna(val) and "Клиент" in str(val):
                            return i
                except Exception:
                    continue
            return None
        except Exception:
            return None

    async def _parse_deals(self, df: pd.DataFrame, period: Period) -> list[Deal]:
        """Parse deals from DataFrame."""
        deals = []
        current_deal = None
        row_counter = 0
        position_counter = 0  # Track position number within current deal

        try:
            # Get column names and handle NaN
            columns = self._normalize_columns(df.columns)
            df.columns = columns

            client_col = columns[0] if len(columns) > 0 else None
            if client_col is None:
                self.stats.warnings.append("No client column found")
                return []

            logger.info(f"Columns: {columns}")

            for _, row in df.iterrows():
                row_counter += 1
                try:
                    client_value = row[client_col]

                    if pd.notna(client_value) and str(client_value).strip():
                        # Master record - new deal
                        if current_deal is not None:
                            # Finalize previous deal (totals calculated automatically via add_item)
                            try:
                                deals.append(current_deal)
                                self.stats.processed_deals += 1
                            except Exception as e:
                                self.stats.failed_deals += 1
                                error_msg = f"Error finalizing deal '{current_deal.client_name}' (row {row_counter}): {str(e)}"
                                self.stats.errors.append(error_msg)

                        # Create new deal and reset position counter
                        position_counter = 0  # Reset position counter for new deal
                        try:
                            current_deal = await self._create_deal_from_row(row, columns, period)
                            self.stats.total_deals += 1
                        except Exception as e:
                            self.stats.failed_deals += 1
                            error_msg = f"Error creating deal from row {row_counter}: {str(e)}"
                            self.stats.errors.append(error_msg)
                            # Create minimal deal to maintain structure
                            current_deal = self._create_error_deal(
                                f"ERROR_ROW_{row_counter}", period
                            )

                    elif current_deal is not None:
                        # Detail record - item
                        try:
                            position_counter += 1  # Increment position number
                            item = await self._create_item_from_row(
                                row, columns, position_counter, current_deal
                            )
                            if item and item.product_name.strip():
                                current_deal.add_item(item)
                                self.stats.processed_items += 1
                            self.stats.total_items += 1
                        except Exception as e:
                            self.stats.failed_items += 1
                            error_msg = f"Error creating item from row {row_counter} for deal '{current_deal.client_name}': {str(e)}"
                            self.stats.errors.append(error_msg)
                            # Add error item to maintain structure
                            position_counter += 1
                            error_item = DealItem(
                                product_name=f"ERROR_ROW_{row_counter}",
                                position_number=position_counter,
                                client_name=current_deal.client_name,
                                period_month=current_deal.period_month,
                                period_year=current_deal.period_year,
                                seller=current_deal.seller,
                                invoice_info=current_deal.invoice_info,
                            )
                            current_deal.add_item(error_item)

                except Exception as e:
                    error_msg = f"Critical error processing row {row_counter}: {str(e)}"
                    self.stats.errors.append(error_msg)
                    logger.error(error_msg)
                    continue

            # Add last deal (totals calculated automatically via add_item)
            if current_deal is not None:
                try:
                    deals.append(current_deal)
                    self.stats.processed_deals += 1
                except Exception as e:
                    self.stats.failed_deals += 1
                    error_msg = f"Error finalizing last deal '{current_deal.client_name}': {str(e)}"
                    self.stats.errors.append(error_msg)
                    # Add anyway
                    deals.append(current_deal)

            return deals

        except Exception as e:
            error_msg = f"Critical error parsing deals: {str(e)}"
            self.stats.errors.append(error_msg)
            logger.error(error_msg)
            return deals

    def _normalize_columns(self, columns) -> list[str]:
        """Normalize column names and handle NaN values."""
        normalized = []
        for i, col in enumerate(columns):
            if pd.notna(col):
                normalized.append(str(col).strip())
            else:
                normalized.append(f"col_{i}")
        return normalized

    async def _create_deal_from_row(
        self, row: pd.Series, columns: list[str], period: Period
    ) -> Deal:
        """Create Deal from master record row using DealBuilder."""

        def safe_get(col_idx: int, default=None):
            try:
                if col_idx < len(columns) and columns[col_idx] in row.index:
                    val = row[columns[col_idx]]
                    return val if pd.notna(val) else default
                return default
            except (KeyError, TypeError):
                return default

        # Basic fields
        client_name = self._safe_string(safe_get(0)) or "UNKNOWN_CLIENT"
        invoice_info_raw = self._safe_string(safe_get(1))
        # Fallback to safe non-empty string if source absent in current layout
        invoice_info = invoice_info_raw or "N/A"
        invoice_number, invoice_date = self._parse_invoice_info(invoice_info_raw)

        # Initialize builder
        builder = DealBuilder(period=period)
        builder.client_name = client_name
        builder.invoice_info = invoice_info
        builder.invoice_number = invoice_number
        builder.invoice_date = invoice_date

        # Optional/status fields
        try:
            if len(columns) > 2:
                shipped_raw = self._safe_string(safe_get(2))
                builder.is_shipped = Status.from_string(shipped_raw)

            if len(columns) > 3:
                builder.upd_number = self._safe_string(safe_get(3))

            if len(columns) > 4:
                paid_raw = self._safe_string(safe_get(4))
                builder.is_paid = Status.from_string(paid_raw)

            if len(columns) > 5:
                revenue_val = safe_get(5)
                if revenue_val is not None:
                    builder.total_revenue = Money(amount=Decimal(str(revenue_val)))

            if len(columns) > 6:
                margin_val = safe_get(6)
                if margin_val is not None:
                    builder.total_margin = SignedMoney(amount=Decimal(str(margin_val)))

            if len(columns) > 7:
                builder.seller = self._safe_string(safe_get(7))

            if len(columns) > 8:
                cost_val = safe_get(8)
                if cost_val is not None:
                    builder.total_cost = Money(amount=Decimal(str(cost_val)))

            if len(columns) > 9:
                kickback_val = safe_get(9)
                if kickback_val is not None:
                    builder.kickback_amount = Money(amount=Decimal(str(kickback_val)))

        except Exception as e:
            warning_msg = f"Warning parsing optional fields for deal '{client_name}': {str(e)}"
            self.stats.warnings.append(warning_msg)
            logger.warning(warning_msg)

        # Build only when everything required is present
        if not builder.is_ready():
            missing_msg = "Deal row missing required fields: "
            if not builder.client_name:
                missing_msg += "client_name "
            if not builder.invoice_info:
                missing_msg += "invoice_info "
            if not builder.seller:
                missing_msg += "seller "
            raise ValueError(missing_msg.strip())

        return builder.build()

    async def _create_item_from_row(self, row: pd.Series, columns: list[str], position_number: int = 1, deal: Deal | None = None) -> DealItem | None:
        """Create DealItem from detail record row with deal context."""

        def safe_get(col_idx: int, default=None):
            try:
                if col_idx < len(columns) and columns[col_idx] in row.index:
                    val = row[columns[col_idx]]
                    return val if pd.notna(val) else default
                return default
            except (KeyError, TypeError):
                return default

        # Check if we have product name
        product_name = self._safe_string(safe_get(1))
        if not product_name:
            return None

        # Provide required context from parent deal
        if deal is None:
            raise ValueError("Deal context is required to create DealItem")

        item = DealItem(
            product_name=product_name,
            position_number=position_number,
            client_name=deal.client_name,
            period_month=deal.period_month,
            period_year=deal.period_year,
            seller=deal.seller,
            invoice_info=deal.invoice_info,
        )

        try:
            # Quantity
            if len(columns) > 2:
                qty_val = safe_get(2)
                if qty_val is not None:
                    item.quantity = Decimal(str(qty_val))

            # Purchase price
            if len(columns) > 3:
                purchase_val = safe_get(3)
                if purchase_val is not None:
                    # Convert to Decimal preserving precision from Excel
                    # If value is already Decimal (preserved from Excel), use it directly
                    # Otherwise convert from float/int/string
                    if isinstance(purchase_val, Decimal):
                        purchase_decimal = purchase_val
                    elif isinstance(purchase_val, (int, float)):
                        # For float values from pandas, convert to Decimal
                        # Note: if pandas converted Decimal to float64, precision may be lost
                        purchase_decimal = Decimal(str(purchase_val))
                    else:
                        purchase_decimal = Decimal(str(purchase_val))
                    item.purchase_price = Money5(amount=purchase_decimal)

            # Sale price
            if len(columns) > 4:
                sale_val = safe_get(4)
                if sale_val is not None:
                    item.sale_price = Money(amount=Decimal(str(sale_val)))

            # Revenue
            if len(columns) > 5:
                revenue_val = safe_get(5)
                if revenue_val is not None:
                    item.revenue = Money(amount=Decimal(str(revenue_val)))

            # Margin
            if len(columns) > 6:
                margin_val = safe_get(6)
                if margin_val is not None:
                    # Convert to Decimal preserving precision from Excel
                    # If value is already Decimal (preserved from Excel), use it directly
                    # Otherwise convert from float/int/string
                    if isinstance(margin_val, Decimal):
                        margin_decimal = margin_val
                    elif isinstance(margin_val, (int, float)):
                        # For float values from pandas, convert to Decimal
                        # Note: if pandas converted Decimal to float64, precision may be lost
                        margin_decimal = Decimal(str(margin_val))
                    else:
                        margin_decimal = Decimal(str(margin_val))
                    item.margin = SignedMoney5(amount=margin_decimal)

            # Cost
            if len(columns) > 8:
                cost_val = safe_get(8)
                if cost_val is not None:
                    item.cost = Money(amount=Decimal(str(cost_val)))

            # Supplier
            if len(columns) > 9:
                item.supplier_name = self._safe_string(safe_get(9))

            # Pickup date
            if len(columns) > 10:
                item.pickup_date = self._safe_pickup_date(safe_get(10))

            # Fields are now calculated automatically via @model_validator

        except Exception as e:
            warning_msg = f"Warning parsing item fields for '{product_name}': {str(e)}"
            self.stats.warnings.append(warning_msg)
            logger.warning(warning_msg)

        return item

    def _create_error_deal(self, client_name: str, period: Period) -> Deal:
        """Create minimal deal for error cases using safe placeholders."""
        builder = DealBuilder(period=period)
        builder.client_name = client_name
        builder.invoice_info = f"ERROR_{client_name}"
        builder.seller = "ERROR"
        return builder.build()

    @staticmethod
    def _safe_string(value) -> str:
        """Safely convert value to string."""
        if pd.isna(value) or value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _safe_pickup_date(value) -> str:
        """Safely convert pickup date with proper number formatting."""
        if pd.isna(value) or value is None:
            return ""

        try:
            num_value = float(value)
            if num_value.is_integer():
                return str(int(num_value))
            else:
                return str(num_value).strip()
        except (ValueError, TypeError):
            return str(value).strip()

    @staticmethod
    def _parse_invoice_info(invoice_info: str) -> tuple[str, str]:
        """Parse invoice information and return number and date."""
        try:
            if not invoice_info:
                return "", ""

            info = str(invoice_info).strip()

            # Pattern: number + "от" + date
            pattern = r"(\d+)\s+от\s+(\d{1,2}\.\d{1,2}\.\d{4})"
            match = re.search(pattern, info)

            if match:
                return match.group(1), match.group(2)

            # Alternative pattern: number + space + date
            pattern2 = r"(\d+)\s+(\d{1,2}\.\d{1,2}\.\d{4})"
            match2 = re.search(pattern2, info)

            if match2:
                return match2.group(1), match2.group(2)

            return "", ""

        except Exception:
            return "", ""

    def _log_final_stats(self) -> None:
        """Log final parsing statistics."""
        logger.info("=" * 80)
        logger.info("PARSING STATISTICS")
        logger.info("=" * 80)

        logger.info(f"📋 SHEETS: {self.stats.processed_sheets}/{self.stats.total_sheets} processed")
        logger.info(f"🏢 DEALS: {self.stats.processed_deals}/{self.stats.total_deals} processed")
        logger.info(f"📦 ITEMS: {self.stats.processed_items}/{self.stats.total_items} processed")

        if self.stats.errors:
            logger.warning(f"❌ ERRORS: {len(self.stats.errors)}")
            for error in self.stats.errors[:5]:  # Show first 5
                logger.warning(f"  - {error}")

        if self.stats.warnings:
            logger.warning(f"⚠️ WARNINGS: {len(self.stats.warnings)}")
            for warning in self.stats.warnings[:3]:  # Show first 3
                logger.warning(f"  - {warning}")

        logger.info(
            f"📊 Success rates: Sheets {self.stats.success_rate_sheets:.1f}%, Deals {self.stats.success_rate_deals:.1f}%, Items {self.stats.success_rate_items:.1f}%"
        )
        logger.info("=" * 80)
