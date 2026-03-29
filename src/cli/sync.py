"""Synchronization commands for the `so-uchet` CLI."""

from __future__ import annotations

import asyncio
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import UUID

import click
from openpyxl import load_workbook

from .common import build_database_manager, configure_logging, run_async

try:
    from application.change_detector import ChangeDetectorService
    from application.excel_parser import ExcelParserService
    from application.excel_parser.models import ParseResult
    from application.sync_orchestrator import SyncConfiguration, SyncOrchestratorService
    from domain.value_objects import Period
    from infrastructure.database.event_store import EventStoreImplementation
    from infrastructure.database.repositories import (
        DealRepositoryImplementation,
        SyncSessionRepositoryImplementation,
    )
    from infrastructure.workers.read_model_builder import ReadModelBuilder
except ImportError:  # pragma: no cover - fallback for repo-local imports
    from src.application.change_detector import ChangeDetectorService
    from src.application.excel_parser import ExcelParserService
    from src.application.excel_parser.models import ParseResult
    from src.application.sync_orchestrator import SyncConfiguration, SyncOrchestratorService
    from src.domain.value_objects import Period
    from src.infrastructure.database.event_store import EventStoreImplementation
    from src.infrastructure.database.repositories import (
        DealRepositoryImplementation,
        SyncSessionRepositoryImplementation,
    )
    from src.infrastructure.workers.read_model_builder import ReadModelBuilder

DEFAULT_EXCEL_FILE = (
    Path(__file__).resolve().parents[2]
    / "data"
    / "real_data_for_testing"
    / "Data_source_excel.xlsx"
)
MONTH_ORDER = {
    "Январь": 1,
    "Февраль": 2,
    "Март": 3,
    "Апрель": 4,
    "Май": 5,
    "Июнь": 6,
    "Июль": 7,
    "Август": 8,
    "Сентябрь": 9,
    "Октябрь": 10,
    "Ноябрь": 11,
    "Декабрь": 12,
}
WORKBOOK_NS = {
    "ss": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}


@dataclass(frozen=True)
class SyncSessionRow:
    """Presentation model for sync session output."""

    session_id: str
    sync_type: str
    status: str
    started_at: str
    finished_at: str
    file_path: str
    total_deals: int
    total_items: int
    errors: int


class FilteredExcelParserService:
    """Thin CLI adapter over the existing parser with period filtering."""

    def __init__(self, parser: ExcelParserService, allowed_periods: list[str] | None = None) -> None:
        self._parser = parser
        self._allowed_periods = set(allowed_periods or [])

    async def parse_file(self, file_path: str, sync_session: object | None) -> ParseResult:
        """Parse Excel and optionally keep only selected periods."""
        allowed_periods = list(self._allowed_periods) if self._allowed_periods else None
        return await self._parser.parse_file(
            file_path,
            sync_session,
            allowed_periods=allowed_periods,
        )


@click.group(help="Операции синхронизации Excel -> DB.")
def sync() -> None:
    """Synchronization command group."""


@sync.command("full")
@click.option(
    "--file",
    "file_path",
    type=click.Path(exists=False, dir_okay=False, path_type=Path),
    default=DEFAULT_EXCEL_FILE,
    show_default=True,
    help="Путь к Excel-файлу.",
)
@click.option(
    "--log-level",
    type=click.Choice(["DEBUG", "INFO"], case_sensitive=False),
    default="DEBUG",
    show_default=True,
    help="Уровень логирования для sync-команды.",
)
@click.pass_context
def sync_full(ctx: click.Context, file_path: Path, log_level: str) -> None:
    """Run full synchronization for all valid periods in the workbook."""
    _override_log_level_if_needed(ctx, log_level)
    result = _execute_sync_command(file_path=file_path, sync_type="full", periods=[])
    _print_sync_summary(result)


@sync.command("partial")
@click.option(
    "--file",
    "file_path",
    type=click.Path(exists=False, dir_okay=False, path_type=Path),
    default=DEFAULT_EXCEL_FILE,
    show_default=True,
    help="Путь к Excel-файлу.",
)
@click.option(
    "--last-periods",
    type=click.IntRange(min=1),
    default=12,
    show_default=True,
    help="Сколько последних периодов включить в синхронизацию.",
)
@click.option(
    "--log-level",
    type=click.Choice(["DEBUG", "INFO"], case_sensitive=False),
    default="DEBUG",
    show_default=True,
    help="Уровень логирования для sync-команды.",
)
@click.pass_context
def sync_partial(ctx: click.Context, file_path: Path, last_periods: int, log_level: str) -> None:
    """Run synchronization for the N latest periods from the workbook."""
    _override_log_level_if_needed(ctx, log_level)
    periods = _select_latest_periods(_validate_excel_path(file_path), limit=last_periods)
    result = _execute_sync_command(file_path=file_path, sync_type="partial", periods=periods)
    _print_sync_summary(result, selected_periods=periods)


@sync.command("periods")
@click.option(
    "--file",
    "file_path",
    type=click.Path(exists=False, dir_okay=False, path_type=Path),
    default=DEFAULT_EXCEL_FILE,
    show_default=True,
    help="Путь к Excel-файлу.",
)
@click.option(
    "--periods",
    multiple=True,
    help='Конкретные периоды, например: --periods "Январь 2025" --periods "Февраль 2025".',
)
@click.option(
    "--log-level",
    type=click.Choice(["DEBUG", "INFO"], case_sensitive=False),
    default="DEBUG",
    show_default=True,
    help="Уровень логирования для sync-команды.",
)
@click.pass_context
def sync_periods(
    ctx: click.Context,
    file_path: Path,
    periods: tuple[str, ...],
    log_level: str,
) -> None:
    """Run synchronization for explicitly selected periods."""
    _override_log_level_if_needed(ctx, log_level)
    validated_file = _validate_excel_path(file_path)
    selected_periods = _resolve_explicit_periods(validated_file, periods)
    result = _execute_sync_command(
        file_path=file_path,
        sync_type="partial",
        periods=selected_periods,
    )
    _print_sync_summary(result, selected_periods=selected_periods)


@sync.command("status")
def sync_status() -> None:
    """Show the running sync session or the latest completed one."""
    row = run_async(_load_sync_status())
    if row is None:
        click.echo("Sync sessions not found")
        return

    click.echo(_format_session_row(row, detailed=True))


@sync.command("list")
@click.option(
    "--limit",
    type=click.IntRange(min=1),
    default=10,
    show_default=True,
    help="Сколько последних сессий показать.",
)
def sync_list(limit: int) -> None:
    """List recent synchronization sessions."""
    rows = run_async(_load_sync_sessions(limit))
    if not rows:
        click.echo("Sync sessions not found")
        return

    click.echo(
        "session_id | type | status | started_at | finished_at | deals | items | errors | file"
    )
    for row in rows:
        click.echo(_format_session_row(row))


def _override_log_level_if_needed(ctx: click.Context, log_level: str) -> None:
    """Apply command-level log level override over the root CLI setting."""
    root_level = (ctx.obj or {}).get("log_level")
    if root_level != log_level.upper():
        configure_logging(log_level.upper())


def _execute_sync_command(file_path: Path, sync_type: str, periods: list[str]) -> Any:
    """Run sync command in async context and convert failures to Click exceptions."""
    try:
        return run_async(
            _run_sync(
                file_path=_validate_excel_path(file_path),
                sync_type=sync_type,
                periods=periods,
            )
        )
    except click.ClickException:
        raise
    except Exception as exc:
        raise click.ClickException(f"Ошибка выполнения sync-команды: {exc}") from exc


async def _run_sync(file_path: Path, sync_type: str, periods: list[str]) -> Any:
    """Execute sync with a single async DB session."""
    db_manager = build_database_manager()
    try:
        async with db_manager.get_async_session() as session:
            orchestrator = _build_sync_orchestrator(session, periods)
            config = SyncConfiguration(
                sync_type=sync_type,
                partial_periods=periods,
                continue_on_errors=False,
                rollback_on_failure=True,
                create_events=True,
                update_read_models=True,
                log_level="DEBUG",
            )
            try:
                result = await orchestrator.execute_sync(str(file_path), config)
                await session.commit()
                return result
            except (KeyboardInterrupt, asyncio.CancelledError):
                if orchestrator.active_session_id is not None:
                    await _force_fail_sync_session(
                        orchestrator.active_session_id,
                        "Marked failed after interrupted CLI sync run",
                    )
                raise
    finally:
        await db_manager.close()


def _build_sync_orchestrator(session: Any, periods: list[str]) -> SyncOrchestratorService:
    """Create the existing sync services bound to one async session."""
    deal_repo = DealRepositoryImplementation(session)
    change_detector = ChangeDetectorService(deal_repo)
    event_store = EventStoreImplementation(session)
    sync_session_repo = SyncSessionRepositoryImplementation(session)
    parser = FilteredExcelParserService(ExcelParserService(), periods if periods else None)
    read_model_builder = ReadModelBuilder(session, event_store)
    return SyncOrchestratorService(
        excel_parser=parser,
        change_detector=change_detector,
        event_store=event_store,
        sync_session_repository=sync_session_repo,
        read_model_builder=read_model_builder,
    )


async def _load_sync_status() -> SyncSessionRow | None:
    """Load current running session or the latest one."""
    db_manager = build_database_manager()
    try:
        async with db_manager.get_async_session() as session:
            repository = SyncSessionRepositoryImplementation(session)
            running = await repository.get_running_session()
            if running is not None:
                return _to_session_row(running)

            latest = await repository.get_latest_sessions(limit=1)
            return _to_session_row(latest[0]) if latest else None
    finally:
        await db_manager.close()


async def _load_sync_sessions(limit: int) -> list[SyncSessionRow]:
    """Load the latest sync sessions."""
    db_manager = build_database_manager()
    try:
        async with db_manager.get_async_session() as session:
            repository = SyncSessionRepositoryImplementation(session)
            sessions = await repository.get_latest_sessions(limit=limit)
            return [_to_session_row(sync_session) for sync_session in sessions]
    finally:
        await db_manager.close()


async def _force_fail_sync_session(session_id: str, error_message: str) -> None:
    """Force-close a pending sync session in a fresh DB transaction."""
    db_manager = build_database_manager()
    try:
        async with db_manager.get_async_session() as session:
            repository = SyncSessionRepositoryImplementation(session)
            sync_session = await repository.get_by_id(UUID(session_id))
            if sync_session is None or sync_session.status.value != "pending":
                await session.rollback()
                return

            sync_session.complete_failed(error_message)
            await repository.save_visible(sync_session)
            await session.rollback()
    finally:
        await db_manager.close()


def _to_session_row(sync_session: Any) -> SyncSessionRow:
    """Convert domain sync session into a presentation row."""
    return SyncSessionRow(
        session_id=str(sync_session.id),
        sync_type=sync_session.sync_type.value,
        status=sync_session.status.value,
        started_at=_format_datetime(sync_session.started_at),
        finished_at=_format_datetime(sync_session.finished_at),
        file_path=sync_session.source_file_path or "-",
        total_deals=sync_session.stats.total_deals,
        total_items=sync_session.stats.total_items,
        errors=len(sync_session.stats.errors),
    )


def _format_datetime(value: Any) -> str:
    """Format timestamps for console output."""
    if value is None:
        return "-"
    return value.isoformat(sep=" ", timespec="seconds")


def _format_session_row(row: SyncSessionRow, detailed: bool = False) -> str:
    """Render sync session information for CLI output."""
    if detailed:
        return (
            f"session_id: {row.session_id}\n"
            f"type: {row.sync_type}\n"
            f"status: {row.status}\n"
            f"started_at: {row.started_at}\n"
            f"finished_at: {row.finished_at}\n"
            f"deals: {row.total_deals}\n"
            f"items: {row.total_items}\n"
            f"errors: {row.errors}\n"
            f"file: {row.file_path}"
        )
    return (
        f"{row.session_id} | {row.sync_type} | {row.status} | {row.started_at} | "
        f"{row.finished_at} | {row.total_deals} | {row.total_items} | {row.errors} | "
        f"{row.file_path}"
    )


def _print_sync_summary(result: Any, selected_periods: list[str] | None = None) -> None:
    """Print a short sync summary to the console."""
    click.echo(f"Sync session: {result.sync_session_id}")
    click.echo(f"Success: {result.summary.success}")
    click.echo(f"File: {result.summary.file_path}")
    if selected_periods:
        click.echo(f"Periods: {', '.join(selected_periods)}")
    click.echo(f"Deals processed: {result.summary.total_deals_processed}")
    click.echo(f"Items processed: {result.summary.total_items_processed}")
    click.echo(
        "Changes: "
        f"+{result.summary.insertions_count} "
        f"~{result.summary.updates_count} "
        f"-{result.summary.deletions_count}"
    )
    click.echo(f"Duration: {result.summary.duration_seconds:.2f}s")
    if result.warnings:
        click.echo(f"Warnings: {len(result.warnings)}")
    if result.errors:
        click.echo(f"Errors: {len(result.errors)}")


def _validate_excel_path(file_path: Path) -> Path:
    """Validate the Excel path before starting sync."""
    resolved = file_path.resolve()
    if not resolved.exists():
        raise click.ClickException(f"Excel file not found: {resolved}")
    if not resolved.is_file():
        raise click.ClickException(f"Path is not a file: {resolved}")
    return resolved


def _select_latest_periods(file_path: Path, limit: int) -> list[str]:
    """Return the latest valid workbook periods in descending order."""
    available_periods = _list_available_periods(file_path)
    selected = available_periods[:limit]
    if not selected:
        raise click.ClickException(f"No valid periods found in Excel file: {file_path}")
    return selected


def _resolve_explicit_periods(file_path: Path, requested_periods: tuple[str, ...]) -> list[str]:
    """Resolve explicit CLI period names against workbook periods."""
    available_periods = _list_available_periods(file_path)
    if not requested_periods:
        return available_periods[:1]

    normalized_available = {period: period for period in available_periods}
    selected: list[str] = []
    missing: list[str] = []

    for raw_period in requested_periods:
        normalized = _normalize_period_name(raw_period)
        if normalized in normalized_available:
            selected.append(normalized_available[normalized])
        else:
            missing.append(raw_period)

    if missing:
        raise click.ClickException(
            "Requested periods not found in Excel file: "
            f"{', '.join(missing)}. Available: {', '.join(available_periods)}"
        )

    return selected


def _list_available_periods(file_path: Path) -> list[str]:
    """List valid period sheet names as canonical 'Month YYYY' values."""
    try:
        sheet_names = _read_sheet_names_from_archive(file_path)
    except Exception:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet_names = list(workbook.sheetnames)
        finally:
            workbook.close()

    canonical_periods: dict[str, tuple[int, int]] = {}
    for sheet_name in sheet_names:
        try:
            period = Period.from_sheet_name(sheet_name)
        except Exception:
            continue
        canonical = f"{period.month} {period.year}"
        canonical_periods[canonical] = _period_sort_key(canonical)

    sorted_periods = sorted(
        canonical_periods.keys(),
        key=lambda value: canonical_periods[value],
        reverse=True,
    )
    return sorted_periods


def _read_sheet_names_from_archive(file_path: Path) -> list[str]:
    """Read workbook sheet names directly from OOXML metadata."""
    with zipfile.ZipFile(file_path) as archive:
        workbook_xml = archive.read("xl/workbook.xml")

    root = ET.fromstring(workbook_xml)
    sheets_element = root.find("ss:sheets", WORKBOOK_NS)
    if sheets_element is None:
        return []

    sheet_names: list[str] = []
    for sheet_element in sheets_element.findall("ss:sheet", WORKBOOK_NS):
        sheet_name = sheet_element.get("name")
        if sheet_name:
            sheet_names.append(sheet_name)

    return sheet_names


def _normalize_period_name(period_name: str) -> str:
    """Normalize arbitrary user input to canonical 'Month YYYY'."""
    stripped = period_name.strip()
    try:
        period = Period.from_sheet_name(stripped)
        return f"{period.month} {period.year}"
    except Exception:
        return stripped


def _period_sort_key(period_name: str) -> tuple[int, int]:
    """Build descending sort key for canonical period names."""
    normalized = _normalize_period_name(period_name)
    month_name, year = normalized.rsplit(" ", 1)
    return int(year), MONTH_ORDER[month_name]
